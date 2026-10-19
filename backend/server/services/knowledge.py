from __future__ import annotations

import hashlib
import json
import mimetypes
import os
import re
import sqlite3
import warnings
from datetime import datetime, timezone
from pathlib import Path
from uuid import uuid4

from server.models.schemas import KnowledgeArtifact, KnowledgeArtifactDetail, KnowledgeSource, KnowledgeSourceCreate, SyncRun

SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx", ".xlsx", ".csv", ".java", ".jsp", ".js", ".ts", ".html", ".css", ".xml", ".json", ".yml", ".yaml", ".properties", ".sql", ".py"}
TEXT_EXTENSIONS = SUPPORTED_EXTENSIONS - {".pdf", ".docx", ".xlsx"}
EXCLUDED_DIRECTORIES = {".git", ".svn", ".idea", ".vscode", "node_modules", "venv", ".venv", "__pycache__", "dist", "build", "target"}
SENSITIVE_SUFFIXES = {".key", ".pem", ".p12", ".pfx"}
SENSITIVE_NAMES = {"id_dsa", "id_ecdsa", "id_ed25519", "id_rsa", "credentials", "credentials.json", "secrets", "secrets.json"}
MAX_FILE_SIZE = 25 * 1024 * 1024
MAX_TEXT_SIZE = 2_000_000
MAX_PDF_PAGES = 500
MAX_XLSX_CELLS = 100_000
CHUNK_TARGET = 1500
CHUNK_MAX = 2000


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _ignored(path: Path) -> bool:
    name = path.name.lower()
    return name == ".env" or name.startswith(".env.") or path.suffix.lower() in SENSITIVE_SUFFIXES or name in SENSITIVE_NAMES or "secret" in name or "private_key" in name


def _hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while data := handle.read(1024 * 1024):
            digest.update(data)
    return digest.hexdigest()


def _chunks(text: str) -> list[tuple[str, int, int]]:
    chunks: list[tuple[str, int, int]] = []
    start = 0
    for block in re.split(r"(?<=\n)\s*\n", text):
        if not block:
            start += len(block)
            continue
        pieces = [block[index:index + CHUNK_MAX] for index in range(0, len(block), CHUNK_MAX)]
        for piece in pieces:
            if chunks and len(chunks[-1][0]) + len(piece) <= CHUNK_MAX and len(chunks[-1][0]) < CHUNK_TARGET:
                previous, previous_start, _ = chunks[-1]
                chunks[-1] = (previous + piece, previous_start, start + len(piece))
            else:
                chunks.append((piece, start, start + len(piece)))
            start += len(piece)
    return [(content, begin, end) for content, begin, end in chunks if content.strip()]


class KnowledgeService:
    """Local-only collection, extraction and archive service backed by WorkMate SQLite."""

    def __init__(self, database_path: str | Path = "data/workmate.db") -> None:
        self.path = Path(database_path)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self._initialize()

    def _connect(self) -> sqlite3.Connection:
        connection = sqlite3.connect(self.path, timeout=5)
        connection.row_factory = sqlite3.Row
        connection.execute("PRAGMA foreign_keys = ON")
        return connection

    def _initialize(self) -> None:
        with self._connect() as db:
            db.executescript("""
                CREATE TABLE IF NOT EXISTS knowledge_sources (
                    id TEXT PRIMARY KEY, name TEXT NOT NULL, local_path TEXT NOT NULL UNIQUE,
                    created_at TEXT NOT NULL, last_synced_at TEXT
                );
                CREATE TABLE IF NOT EXISTS knowledge_artifacts (
                    id TEXT PRIMARY KEY, source_id TEXT NOT NULL REFERENCES knowledge_sources(id) ON DELETE CASCADE,
                    relative_path TEXT NOT NULL, title TEXT NOT NULL, mime_type TEXT, sha256 TEXT NOT NULL,
                    size_bytes INTEGER NOT NULL, modified_at TEXT NOT NULL, is_deleted INTEGER NOT NULL DEFAULT 0,
                    created_at TEXT NOT NULL, updated_at TEXT NOT NULL, UNIQUE(source_id, relative_path)
                );
                CREATE TABLE IF NOT EXISTS knowledge_documents (
                    id TEXT PRIMARY KEY, artifact_id TEXT NOT NULL UNIQUE REFERENCES knowledge_artifacts(id) ON DELETE CASCADE,
                    content TEXT NOT NULL, parser_type TEXT NOT NULL, created_at TEXT NOT NULL, updated_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS knowledge_chunks (
                    id TEXT PRIMARY KEY, document_id TEXT NOT NULL REFERENCES knowledge_documents(id) ON DELETE CASCADE,
                    chunk_index INTEGER NOT NULL, content TEXT NOT NULL, character_start INTEGER NOT NULL,
                    character_end INTEGER NOT NULL, token_count INTEGER NOT NULL, sha256 TEXT NOT NULL,
                    UNIQUE(document_id, chunk_index)
                );
                CREATE TABLE IF NOT EXISTS sync_runs (
                    id TEXT PRIMARY KEY, source_id TEXT NOT NULL REFERENCES knowledge_sources(id) ON DELETE CASCADE,
                    status TEXT NOT NULL, total INTEGER NOT NULL DEFAULT 0, scanned INTEGER NOT NULL DEFAULT 0, created INTEGER NOT NULL DEFAULT 0,
                    updated INTEGER NOT NULL DEFAULT 0, unchanged INTEGER NOT NULL DEFAULT 0, deleted INTEGER NOT NULL DEFAULT 0,
                    parsed INTEGER NOT NULL DEFAULT 0, failed INTEGER NOT NULL DEFAULT 0, chunks_created INTEGER NOT NULL DEFAULT 0,
                    error TEXT, started_at TEXT NOT NULL, finished_at TEXT
                );
            """)
            columns = {row["name"] for row in db.execute("PRAGMA table_info(sync_runs)")}
            if "cancel_requested" not in columns:
                db.execute("ALTER TABLE sync_runs ADD COLUMN cancel_requested INTEGER NOT NULL DEFAULT 0")
            if "total" not in columns:
                db.execute("ALTER TABLE sync_runs ADD COLUMN total INTEGER NOT NULL DEFAULT 0")

    def list_sources(self) -> list[KnowledgeSource]:
        with self._connect() as db:
            rows = db.execute("SELECT s.*, COUNT(a.id) artifact_count FROM knowledge_sources s LEFT JOIN knowledge_artifacts a ON a.source_id=s.id AND a.is_deleted=0 GROUP BY s.id ORDER BY s.name").fetchall()
        return [KnowledgeSource(id=row["id"], name=row["name"], local_path=row["local_path"], created_at=row["created_at"], last_synced_at=row["last_synced_at"], artifact_count=row["artifact_count"]) for row in rows]

    def create_source(self, payload: KnowledgeSourceCreate) -> KnowledgeSource:
        path = Path(payload.local_path).expanduser()
        if not path.exists() or not path.is_dir():
            raise ValueError("접근 가능한 로컬 폴더 경로를 입력해 주세요.")
        if not os.access(path, os.R_OK | os.X_OK):
            raise ValueError("등록한 폴더를 읽을 권한이 없습니다.")
        source_id, now = str(uuid4()), _now()
        try:
            with self._connect() as db:
                db.execute("INSERT INTO knowledge_sources VALUES (?, ?, ?, ?, NULL)", (source_id, payload.name.strip() or path.name, str(path.resolve()), now))
        except sqlite3.IntegrityError as error:
            raise ValueError("이미 등록된 폴더입니다.") from error
        return KnowledgeSource(id=source_id, name=payload.name.strip() or path.name, local_path=str(path.resolve()), created_at=now, artifact_count=0)

    def start_scan(self, source_id: str) -> SyncRun:
        with self._connect() as db:
            if not db.execute("SELECT 1 FROM knowledge_sources WHERE id=?", (source_id,)).fetchone(): raise KeyError("지식 폴더를 찾을 수 없습니다.")
            run_id, started = str(uuid4()), _now()
            db.execute("INSERT INTO sync_runs (id,source_id,status,started_at) VALUES (?,?,?,?)", (run_id, source_id, "running", started))
            return self._run(db.execute("SELECT * FROM sync_runs WHERE id=?", (run_id,)).fetchone())

    def cancel(self, run_id: str) -> SyncRun:
        with self._connect() as db:
            row = db.execute("SELECT * FROM sync_runs WHERE id=?", (run_id,)).fetchone()
            if not row: raise KeyError("수집 작업을 찾을 수 없습니다.")
            if row["status"] == "running": db.execute("UPDATE sync_runs SET cancel_requested=1 WHERE id=?", (run_id,))
            return self._run(db.execute("SELECT * FROM sync_runs WHERE id=?", (run_id,)).fetchone())

    def get_run(self, run_id: str) -> SyncRun:
        with self._connect() as db:
            row = db.execute("SELECT * FROM sync_runs WHERE id=?", (run_id,)).fetchone()
        if not row: raise KeyError("수집 작업을 찾을 수 없습니다.")
        return self._run(row)

    def scan(self, source_id: str, run_id: str | None = None) -> SyncRun:
        with self._connect() as db:
            source = db.execute("SELECT * FROM knowledge_sources WHERE id=?", (source_id,)).fetchone()
            if not source:
                raise KeyError("지식 폴더를 찾을 수 없습니다.")
            root = Path(source["local_path"])
            if not root.is_dir() or not os.access(root, os.R_OK | os.X_OK):
                raise ValueError("등록된 폴더에 접근할 수 없습니다.")
            if run_id is None:
                run_id, started = str(uuid4()), _now()
                db.execute("INSERT INTO sync_runs (id,source_id,status,started_at) VALUES (?,?,?,?)", (run_id, source_id, "running", started))
            result = {"scanned": 0, "created": 0, "updated": 0, "unchanged": 0, "deleted": 0, "parsed": 0, "failed": 0, "chunks_created": 0}
            seen: set[str] = set()
            try:
                candidates: list[Path] = []
                for directory, directories, files in os.walk(root, topdown=True, followlinks=False):
                    directories[:] = [name for name in directories if name.lower() not in EXCLUDED_DIRECTORIES]
                    candidates.extend(path for filename in files if (path := Path(directory, filename)).suffix.lower() in SUPPORTED_EXTENSIONS and not path.is_symlink() and not _ignored(path))
                db.execute("UPDATE sync_runs SET total=? WHERE id=?", (len(candidates), run_id))
                db.commit()
                for path in candidates:
                        if db.execute("SELECT cancel_requested FROM sync_runs WHERE id=?", (run_id,)).fetchone()[0]:
                            finished = _now(); db.execute("UPDATE sync_runs SET status='cancelled',finished_at=?,scanned=?,created=?,updated=?,unchanged=?,deleted=?,parsed=?,failed=?,chunks_created=? WHERE id=?", (finished, result["scanned"], result["created"], result["updated"], result["unchanged"], result["deleted"], result["parsed"], result["failed"], result["chunks_created"], run_id)); return self._run(db.execute("SELECT * FROM sync_runs WHERE id=?", (run_id,)).fetchone())
                        relative = path.relative_to(root).as_posix(); seen.add(relative); result["scanned"] += 1
                        db.execute("UPDATE sync_runs SET scanned=? WHERE id=?", (result["scanned"], run_id))
                        stat = path.stat(); modified = datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat()
                        old = db.execute("SELECT * FROM knowledge_artifacts WHERE source_id=? AND relative_path=?", (source_id, relative)).fetchone()
                        if old and not old["is_deleted"] and old["size_bytes"] == stat.st_size and old["modified_at"] == modified:
                            result["unchanged"] += 1; db.commit(); continue
                        digest = _hash(path)
                        if old is None:
                            artifact_id = str(uuid4()); now = _now()
                            db.execute("INSERT INTO knowledge_artifacts VALUES (?,?,?,?,?,?,?,?,?,?,?)", (artifact_id, source_id, relative, path.name, mimetypes.guess_type(path.name)[0], digest, stat.st_size, modified, 0, now, now,))
                            result["created"] += 1
                        else:
                            artifact_id = old["id"]
                            if old["sha256"] == digest and not old["is_deleted"]:
                                db.execute("UPDATE knowledge_artifacts SET title=?,mime_type=?,size_bytes=?,modified_at=?,updated_at=? WHERE id=?", (path.name, mimetypes.guess_type(path.name)[0], stat.st_size, modified, _now(), artifact_id)); result["unchanged"] += 1; db.commit(); continue
                            db.execute("UPDATE knowledge_artifacts SET title=?,mime_type=?,sha256=?,size_bytes=?,modified_at=?,is_deleted=0,updated_at=? WHERE id=?", (path.name, mimetypes.guess_type(path.name)[0], digest, stat.st_size, modified, _now(), artifact_id)); db.execute("DELETE FROM knowledge_documents WHERE artifact_id=?", (artifact_id,)); result["updated"] += 1
                        try:
                            content, parser_type = self._extract(path)
                            document_id = str(uuid4()); now = _now()
                            db.execute("INSERT INTO knowledge_documents VALUES (?,?,?,?,?,?)", (document_id, artifact_id, content, parser_type, now, now))
                            chunks = _chunks(content)
                            db.executemany("INSERT INTO knowledge_chunks VALUES (?,?,?,?,?,?,?,?)", [(str(uuid4()), document_id, index, part, begin, end, len(re.findall(r"\\S+", part)), hashlib.sha256(part.encode()).hexdigest()) for index, (part, begin, end) in enumerate(chunks)])
                            result["parsed"] += 1; result["chunks_created"] += len(chunks)
                        except Exception:
                            result["failed"] += 1
                        db.commit()
                for row in db.execute("SELECT id,relative_path FROM knowledge_artifacts WHERE source_id=? AND is_deleted=0", (source_id,)).fetchall():
                    if row["relative_path"] not in seen:
                        db.execute("UPDATE knowledge_artifacts SET is_deleted=1,updated_at=? WHERE id=?", (_now(), row["id"])); db.execute("DELETE FROM knowledge_documents WHERE artifact_id=?", (row["id"],)); result["deleted"] += 1
                finished = _now(); db.execute("UPDATE knowledge_sources SET last_synced_at=? WHERE id=?", (finished, source_id))
                db.execute("UPDATE sync_runs SET status='completed', finished_at=?, scanned=?, created=?, updated=?, unchanged=?, deleted=?, parsed=?, failed=?, chunks_created=? WHERE id=?", (finished, result["scanned"], result["created"], result["updated"], result["unchanged"], result["deleted"], result["parsed"], result["failed"], result["chunks_created"], run_id))
            except Exception as error:
                db.execute("UPDATE sync_runs SET status='failed', finished_at=?, error=? WHERE id=?", (_now(), str(error), run_id))
                raise
            row = db.execute("SELECT * FROM sync_runs WHERE id=?", (run_id,)).fetchone()
        return self._run(row)

    def _extract(self, path: Path) -> tuple[str, str]:
        if path.stat().st_size > MAX_FILE_SIZE: raise ValueError("파일이 25 MiB 제한을 초과했습니다.")
        ext = path.suffix.lower()
        if ext in TEXT_EXTENSIONS:
            raw = path.read_bytes(); text = next((raw.decode(codec) for codec in ("utf-8-sig", "utf-8", "cp949", "euc-kr") if self._can_decode(raw, codec)), raw.decode("utf-8", errors="replace")); parser = "text"
        elif ext == ".pdf":
            from pypdf import PdfReader
            with warnings.catch_warnings():
                warnings.filterwarnings("ignore", message="fontTools is required.*")
                reader = PdfReader(str(path))
            if len(reader.pages) > MAX_PDF_PAGES: raise ValueError("PDF 페이지 수가 500페이지 제한을 초과했습니다.")
            text = "\n\n".join(f"[Page: {index}]\n{page.extract_text() or ''}" for index, page in enumerate(reader.pages, 1)); parser = "pdf"
        elif ext == ".docx":
            from docx import Document
            document = Document(str(path)); text = "\n\n".join([paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()] + ["\n".join(" | ".join(cell.text.replace("\n", " ").strip() for cell in row.cells) for row in table.rows) for table in document.tables]); parser = "docx"
        else:
            from openpyxl import load_workbook
            with warnings.catch_warnings():
                warnings.filterwarnings("ignore", message="Print area cannot be set.*")
                book = load_workbook(str(path), read_only=True, data_only=True)
            cells = 0; pages = []
            try:
                for sheet in book.worksheets:
                    lines = [f"[Sheet: {sheet.title}]"]
                    for row in sheet.iter_rows(values_only=True):
                        cells += len(row)
                        if cells > MAX_XLSX_CELLS: raise ValueError("XLSX 셀 수가 100,000개 제한을 초과했습니다.")
                        values = ["" if value is None else str(value).replace("\n", " ") for value in row]
                        if any(values): lines.append(" | ".join(values))
                    if len(lines) > 1: pages.append("\n".join(lines))
            finally: book.close()
            text, parser = "\n\n".join(pages), "xlsx"
        if not text.strip(): raise ValueError("추출 가능한 텍스트가 없습니다.")
        if len(text) > MAX_TEXT_SIZE: raise ValueError("추출 텍스트가 안전 한도를 초과했습니다.")
        return text, parser

    @staticmethod
    def _can_decode(value: bytes, codec: str) -> bool:
        try: value.decode(codec); return True
        except UnicodeDecodeError: return False

    def list_artifacts(self, source_id: str) -> list[KnowledgeArtifact]:
        with self._connect() as db:
            rows = db.execute("SELECT a.*, COUNT(c.id) chunk_count FROM knowledge_artifacts a LEFT JOIN knowledge_documents d ON d.artifact_id=a.id LEFT JOIN knowledge_chunks c ON c.document_id=d.id WHERE a.source_id=? AND a.is_deleted=0 GROUP BY a.id ORDER BY a.relative_path", (source_id,)).fetchall()
        return [KnowledgeArtifact(id=row["id"], source_id=row["source_id"], title=row["title"], relative_path=row["relative_path"], size_bytes=row["size_bytes"], modified_at=row["modified_at"], chunk_count=row["chunk_count"], parsed=bool(row["chunk_count"])) for row in rows]

    def artifact(self, artifact_id: str) -> KnowledgeArtifactDetail:
        with self._connect() as db:
            row = db.execute("SELECT a.*, d.content, COUNT(c.id) chunk_count FROM knowledge_artifacts a LEFT JOIN knowledge_documents d ON d.artifact_id=a.id LEFT JOIN knowledge_chunks c ON c.document_id=d.id WHERE a.id=? AND a.is_deleted=0 GROUP BY a.id", (artifact_id,)).fetchone()
        if not row: raise KeyError("문서를 찾을 수 없습니다.")
        return KnowledgeArtifactDetail(id=row["id"], source_id=row["source_id"], title=row["title"], relative_path=row["relative_path"], size_bytes=row["size_bytes"], modified_at=row["modified_at"], chunk_count=row["chunk_count"], parsed=bool(row["chunk_count"]), content=row["content"] or "")

    @staticmethod
    def _run(row: sqlite3.Row) -> SyncRun:
        return SyncRun(**dict(row))
